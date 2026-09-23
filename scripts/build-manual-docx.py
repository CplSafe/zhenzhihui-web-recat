# -*- coding: utf-8 -*-
"""把 docs 下的手册 Markdown 转成内嵌图片的 .docx，供飞书「导入本地文件」使用。
图片压到 1200px 宽 JPEG，保证整份文档 < 10MB。

用法：
  python scripts/build-manual-docx.py                         # 主手册 → docs/帧智汇使用手册.docx
  python scripts/build-manual-docx.py docs/帧智汇邀请返利手册.md  # 指定源 md，输出同名 .docx
  python scripts/build-manual-docx.py <src.md> <out.docx>     # 显式指定输出路径
"""
import io
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(r"F:\项目\zhenzhihui-web-recat\docs")
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "帧智汇使用手册.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else SRC.with_suffix(".docx")
IMG_W = 1200
JPEG_Q = 78

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
style.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.9)

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(p, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(0x5A, 0x3E, 0xC8)
        else:
            p.add_run(part)


def add_image(rel):
    src = ROOT / rel
    if not src.exists():
        doc.add_paragraph(f"[缺图: {rel}]")
        return
    im = Image.open(src).convert("RGB")
    if im.width > IMG_W:
        im = im.resize((IMG_W, round(im.height * IMG_W / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=JPEG_Q, optimize=True)
    buf.seek(0)
    doc.add_picture(buf, width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if not (i == 1 and all(set(c) <= set(":- ") for c in r))]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            cell.text = ""
            add_runs(cell.paragraphs[0], c)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s or s == "---":
        i += 1
        continue
    m = re.match(r"^(#{1,6})\s+(.*)", s)
    if m:
        level = len(m.group(1))
        doc.add_heading(m.group(2), level=min(level, 4))
        i += 1
        continue
    m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", s)
    if m:
        add_image(m.group(1))
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(lines[i])
            i += 1
        add_table(rows)
        continue
    if s.startswith(">"):
        body = s.lstrip("> ").strip()
        if not body:
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        if body.startswith("▶"):
            # 视频占位：醒目底色，提醒导入飞书后在此处插入视频
            add_runs(p, body + "　　← 导入飞书后在此处插入该视频")
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xB4, 0x1E, 0x1E)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "FFF1C2")
            pPr.append(shd)
        else:
            add_runs(p, body)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        i += 1
        continue
    m = re.match(r"^(\d+)\.\s+(.*)", s)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, m.group(2))
        i += 1
        continue
    if s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, s[2:])
        i += 1
        continue
    p = doc.add_paragraph()
    add_runs(p, s)
    i += 1

doc.save(OUT)
print("saved", OUT, round(OUT.stat().st_size / 1024 / 1024, 2), "MB")
